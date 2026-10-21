import os
import re
import uuid

import docx
import httpx
import pptx
from bs4 import BeautifulSoup


def save_generated_document(user_id: int, title: str, sections: list[dict]) -> tuple[str, str]:
    d = docx.Document()
    d.add_heading(title, 0)
    for s in sections:
        d.add_heading(s.get("heading", ""), level=1)
        d.add_paragraph(s.get("text", ""))
    name, path = _write(user_id, title, ".docx")
    d.save(path)
    return name, path


def save_generated_presentation(user_id: int, title: str, slides: list[dict]) -> tuple[str, str]:
    p = pptx.Presentation()
    p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = title
    for s in slides:
        slide = p.slides.add_slide(p.slide_layouts[1])
        slide.shapes.title.text = s.get("title", "")
        body = slide.placeholders[1].text_frame
        first = True
        for b in s.get("bullets", [])[:12]:
            para = body.paragraphs[0] if first else body.add_paragraph()
            first = False
            para.text = b
    name, path = _write(user_id, title, ".pptx")
    p.save(path)
    return name, path


def read_link(url: str) -> str:
    r = httpx.get(url, timeout=30, follow_redirects=True, headers={"User-Agent": "Mozilla/5.0 WebHub"})
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form"]):
        tag.decompose()
    text = re.sub(r"\n\s*\n+", "\n\n", soup.get_text("\n", strip=True))
    return text[:8000]


def _write(user_id: int, title: str, ext: str) -> tuple[str, str]:
    safe = "".join(c for c in title if c.isalnum() or c in " ._-" ).strip() or "file"
    name = f"{safe[:80]}{ext}"
    d = os.path.join("/app/data/uploads", str(user_id))
    os.makedirs(d, exist_ok=True)
    path = os.path.join(d, uuid.uuid4().hex[:8] + "_" + name)
    return name, path
